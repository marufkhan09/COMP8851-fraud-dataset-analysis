#!/usr/bin/env python3
import json, sys, subprocess, time, hashlib
from pathlib import Path

# Self-contained reporting dependency setup. Does not alter BWGNN model code/config.
try:
    from docx import Document
    from docx.shared import Inches, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
except ImportError:
    print('python-docx missing; installing reporting-only dependency...', flush=True)
    subprocess.check_call([sys.executable, '-m', 'pip', 'install', '--no-cache-dir', 'python-docx'])
    from docx import Document
    from docx.shared import Inches, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT

WORK=Path('/workspace/bwgnn_vast')
OUT=WORK/'reports/BWGNN_Vast_Progress_Report.docx'
OUT.parent.mkdir(parents=True, exist_ok=True)
DATASETS=[('YelpChi','yelpchi'),('Amazon','amazon'),('T-Finance','tfinance'),('FDCompCN','fdcompcn'),('Elliptic','elliptic'),('T-Social','tsocial')]
RATIOS=['TR40','TR30','TR20','TR10']

def fmt(x, n=4):
    if x is None: return '—'
    return f'{x:.{n}f}'

def ms(x):
    if not x: return '—'
    return f"{x['mean']:.4f} ± {x['sd']:.4f}"

def sha256(path):
    h=hashlib.sha256()
    with open(path,'rb') as f:
        for b in iter(lambda:f.read(8*1024*1024),b''): h.update(b)
    return h.hexdigest()

records=[]
for display,slug in DATASETS:
    summary_path=WORK/f'results/bwgnn/{slug}/final/summary.json'
    manifest_path=WORK/f'evidence/{slug}/input_manifest.json'
    if summary_path.exists():
        s=json.loads(summary_path.read_text())
        m=json.loads(manifest_path.read_text()) if manifest_path.exists() else {}
        records.append({'display':display,'slug':slug,'status':'COMPLETE','summary':s,'manifest':m})
    else:
        smoke=WORK/f'evidence/{slug}/smoke.json'
        status='SMOKE_ONLY' if smoke.exists() else 'NOT_RUN_ON_VAST'
        records.append({'display':display,'slug':slug,'status':status,'summary':None,'manifest':{}})

complete=sum(r['status']=='COMPLETE' for r in records)
doc=Document()
sec=doc.sections[0]
sec.top_margin=Inches(0.55); sec.bottom_margin=Inches(0.55); sec.left_margin=Inches(0.6); sec.right_margin=Inches(0.6)
styles=doc.styles
styles['Normal'].font.name='Arial'; styles['Normal'].font.size=Pt(9)
styles['Title'].font.name='Arial'; styles['Title'].font.size=Pt(18)
styles['Heading 1'].font.name='Arial'; styles['Heading 1'].font.size=Pt(13)
styles['Heading 2'].font.name='Arial'; styles['Heading 2'].font.size=Pt(11)

t=doc.add_paragraph(); t.style='Title'; t.alignment=WD_ALIGN_PARAGRAPH.CENTER; t.add_run('COMP8851 — BWGNN Vast.ai Progress Report')
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run(f'Generated {time.strftime("%Y-%m-%d %H:%M UTC", time.gmtime())} | Completed datasets: {complete}/6').bold=True

doc.add_heading('1. Current coverage', level=1)
table=doc.add_table(rows=1, cols=4); table.alignment=WD_TABLE_ALIGNMENT.CENTER; table.style='Table Grid'
for i,h in enumerate(['Dataset','Vast status','Final runs','Evidence archive']): table.rows[0].cells[i].text=h
for r in records:
    row=table.add_row().cells
    row[0].text=r['display']; row[1].text=r['status']
    row[2].text=str(r['summary'].get('completed_runs','—')) if r['summary'] else '—'
    arc=WORK/f"archives/bwgnn_{r['slug']}_vast_a6000_final.tar.gz"
    row[3].text=arc.name if arc.exists() else '—'

doc.add_paragraph('Coverage is based only on completed Vast.ai controlled result files currently present in the workspace; prior Kaggle/T4 pilot evidence is not counted as a completed Vast dataset.')

for r in records:
    if r['status']!='COMPLETE': continue
    s=r['summary']; m=r['manifest']
    doc.add_heading(f"2.{records.index(r)+1} {r['display']} — COMPLETE", level=2)
    cfg=m.get('frozen_config',{})
    p=doc.add_paragraph()
    p.add_run('Configuration: ').bold=True
    p.add_run(f"hidden={cfg.get('hidden','—')}, order={cfg.get('order','—')}, lr={cfg.get('lr','—')}, weight_decay={cfg.get('weight_decay','—')}; split seed={m.get('split_seed','—')}; train seeds={m.get('train_seeds','—')}.")
    p=doc.add_paragraph()
    p.add_run('Environment: ').bold=True
    p.add_run(f"GPU={m.get('gpu','—')}; Torch={m.get('torch','—')}; DGL={m.get('dgl','—')}; repository commit={m.get('repository_commit','—')}.")
    p=doc.add_paragraph()
    p.add_run('Dataset SHA256: ').bold=True; p.add_run(str(m.get('dataset_sha256','—')))
    tab=doc.add_table(rows=1, cols=7); tab.style='Table Grid'; tab.alignment=WD_TABLE_ALIGNMENT.CENTER
    heads=['Ratio','AUPRC','AUROC','Macro-F1','Fraud recall','Train s/epoch','Peak GPU MB']
    for i,h in enumerate(heads): tab.rows[0].cells[i].text=h
    for ratio in RATIOS:
        rr=s.get('ratios',{}).get(ratio,{})
        row=tab.add_row().cells
        vals=[ratio,ms(rr.get('auprc')),ms(rr.get('auroc')),ms(rr.get('macro_f1')),ms(rr.get('fraud_recall')),ms(rr.get('mean_epoch_train_seconds')),ms(rr.get('peak_gpu_memory_mb'))]
        for i,v in enumerate(vals): row[i].text=v
    arc=WORK/f"archives/bwgnn_{r['slug']}_vast_a6000_final.tar.gz"
    if arc.exists():
        p=doc.add_paragraph(); p.add_run('Evidence archive: ').bold=True; p.add_run(str(arc)); p.add_run('\nSHA256: ').bold=True; p.add_run(sha256(arc))

doc.add_heading('3. Protocol notes', level=1)
for text in [
    'Final runs use TR40/TR30/TR20/TR10 with training seeds 2, 42 and 72.',
    'Checkpoint selection is validation AUPRC; decision threshold is selected from validation Macro-F1 with the frozen tie-break rule.',
    'Test evaluation occurs only after configuration, checkpoint and validation threshold are frozen.',
    'Training time and validation time are recorded separately; failure/OOM status remains part of the benchmark evidence.'
]:
    doc.add_paragraph(text, style='List Bullet')

doc.save(OUT)
print(f'REPORT_COMPLETE={OUT}')
print(f'COMPLETED_DATASETS={complete}/6')
